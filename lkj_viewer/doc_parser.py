# -*- coding: utf-8 -*-
"""文档解析器
从 PDF / Word / Excel 文档中提取 LKJ 基硎数据，按工作簿归类。
返回格式: {sheet_name: {"headers": [...], "rows": [[...], ...]}}
"""

import os
import re


SHEET_MATCH_RULES = [
    ("12.起讫里程", ["正线起讫"]),
    ("9.隧道表", ["隧道号"]),
    ("8.桥梁表", ["桥号", "桥名"]),
    ("7.曲线表", ["曲线方向", "曲线半径"]),
    ("6.坡道表", ["坡度", "坡长"]),
    ("5.线路允许速度表", ["速度区段"]),
    ("4.道岔表", ["道岔编号", "辙叉号"]),
    ("3.股道表", ["股道编号", "有效长"]),
    ("2.车站表", ["车站名", "车站编号", "股道数"]),
    ("1.线路名称表", ["局名", "线名", "线编号"]),
]


def clean_cell(val):
    """清理单元格数据"""
    if val is None:
        return None
    s = str(val).replace("\n", "").replace("\r", "").strip()
    if s == "":
        return None
    try:
        i = int(s)
        return i
    except (ValueError, TypeError):
        pass
    try:
        f = float(s)
        return f
    except (ValueError, TypeError):
        pass
    return s


def clean_row(row):
    """清理一行数据，返回清理后的列表"""
    return [clean_cell(c) for c in row]


def match_sheet(headers):
    """根据列头匹配对应的工作簿名称（子串匹配）"""
    cleaned = [str(h).replace("\n", "").strip() if h else "" for h in headers]
    for sheet_name, keywords in SHEET_MATCH_RULES:
        if all(any(kw in h for h in cleaned) for kw in keywords):
            return sheet_name
    return None


def is_header_row(row):
    """判断一行是否为列头行（而非数据行）"""
    if not row:
        return False
    text = " ".join(str(c).replace("\n", "").strip() if c else "" for c in row)
    header_keywords = ["局名", "线名", "车站名", "桥号", "隧道号",
                       "起点里程", "曲线方向", "坡度", "速度区段",
                       "正线起讫", "道岔编号", "股道编号", "行别"]
    matches = sum(1 for kw in header_keywords if kw in text)
    return matches >= 2


def _fix_line_names(parsed_data):
    """后处理：用线路名称表中的正确线名修正其他表中的残缺线名"""
    if "1.线路名称表" not in parsed_data:
        return
    valid_names = set()
    for row in parsed_data["1.线路名称表"]["rows"]:
        if len(row) > 2 and row[2]:
            valid_names.add(str(row[2]))
    if not valid_names:
        return
    for sheet_name, data in parsed_data.items():
        if sheet_name == "1.线路名称表":
            continue
        headers = data["headers"]
        if "线名" not in headers:
            continue
        name_col = headers.index("线名")
        for row in data["rows"]:
            if name_col >= len(row) or not row[name_col]:
                continue
            current = str(row[name_col]).strip()
            if current in valid_names:
                continue
            best = None
            for vn in valid_names:
                if current == vn:
                    best = vn
                    break
                if current in vn or vn in current:
                    best = vn
                    break
            if best is None:
                for vn in valid_names:
                    trimmed = current
                    while trimmed.startswith("路") or trimmed.startswith("铁"):
                        trimmed = trimmed[1:]
                    if trimmed and (trimmed in vn or vn in trimmed):
                        best = vn
                        break
            if best:
                row[name_col] = best


def parse_pdf(filepath):
    """解析 PDF 文件，返回 {sheet_name: {"headers": [...], "rows": [...]}}"""
    import pdfplumber

    result = {}
    pdf = pdfplumber.open(filepath)

    for page in pdf.pages:
        tables = page.extract_tables()
        for table in tables:
            if not table or len(table) < 2:
                continue
            header = table[0]
            sheet_name = match_sheet(header)
            if sheet_name is None:
                continue

            clean_headers = [str(h).replace("\n", "").strip() if h else f"列{i+1}"
                             for i, h in enumerate(header)]

            if sheet_name not in result:
                result[sheet_name] = {"headers": clean_headers, "rows": []}

            for row in table[1:]:
                if is_header_row(row):
                    continue
                cleaned = clean_row(row)
                if all(c is None for c in cleaned):
                    continue
                result[sheet_name]["rows"].append(cleaned)

    pdf.close()
    return result


def parse_docx(filepath):
    """解析 Word 文档，返回 {sheet_name: {"headers": [...], "rows": [...]}}"""
    import docx

    result = {}
    doc = docx.Document(filepath)

    for table in doc.tables:
        if len(table.rows) < 2:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        sheet_name = match_sheet(header)
        if sheet_name is None:
            continue

        clean_headers = [h if h else f"列{i+1}" for i, h in enumerate(header)]

        if sheet_name not in result:
            result[sheet_name] = {"headers": clean_headers, "rows": []}

        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            if is_header_row(values):
                continue
            cleaned = clean_row(values)
            if all(c is None for c in cleaned):
                continue
            result[sheet_name]["rows"].append(cleaned)

    return result


def parse_excel(filepath):
    """解析 Excel 文件，返回 {sheet_name: {"headers": [...], "rows": [...]}}"""
    import openpyxl

    result = {}
    wb = openpyxl.load_workbook(filepath, data_only=True)

    for ws_name in wb.sheetnames:
        ws = wb[ws_name]
        if ws.max_row < 3:
            continue

        header_row = None
        header_idx = None
        for r in range(1, min(ws.max_row, 6) + 1):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            if is_header_row(row_vals):
                header_row = row_vals
                header_idx = r
                break

        if header_row is None:
            continue

        max_col = 1
        for val in header_row:
            if val is not None:
                idx = header_row.index(val) + 1
                if idx > max_col:
                    max_col = idx

        clean_headers = [str(h).strip() if h else f"列{i+1}"
                         for i, h in enumerate(header_row[:max_col])]

        sheet_name = match_sheet(header_row[:max_col])
        if sheet_name is None:
            sheet_name = ws_name

        rows = []
        for r in range(header_idx + 1, ws.max_row + 1):
            row_data = []
            has_data = False
            for c in range(1, max_col + 1):
                val = ws.cell(row=r, column=c).value
                if val is not None and str(val).strip():
                    has_data = True
                row_data.append(val)
            if has_data:
                rows.append(clean_row(row_data))

        if sheet_name not in result:
            result[sheet_name] = {"headers": clean_headers, "rows": []}
        result[sheet_name]["rows"].extend(rows)

    wb.close()
    return result


def parse_document(filepath):
    """根据文件扩展名自动选择解析器"""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        result = parse_pdf(filepath)
    elif ext in (".docx", ".doc"):
        result = parse_docx(filepath)
    elif ext in (".xlsx", ".xls"):
        result = parse_excel(filepath)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")
    _fix_line_names(result)
    return result


def get_summary(parsed_data):
    """返回解析结果的概要信息"""
    lines = []
    for sheet_name, data in parsed_data.items():
        lines.append(f"  {sheet_name}: {len(data['rows'])} 行, {len(data['headers'])} 列")
    return "\n".join(lines)