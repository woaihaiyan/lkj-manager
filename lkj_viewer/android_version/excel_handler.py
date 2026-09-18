# -*- coding: utf-8 -*-
"""Excel 文件读写处理模块
负责加载 xlsx 各工作簿数据、保存修改回原 Excel。
表结构: 第1行标题, 第2行表编号(LKJ数-X), 第3行列头, 第4行起为数据。
"""

import openpyxl


HEADER_ROW = 3
DATA_START_ROW = 4
TITLE_ROW = 1
CODE_ROW = 2


class SheetData:
    """单个工作簿的数据容器"""

    def __init__(self, name, title, code, headers, rows, max_col):
        self.name = name
        self.title = title
        self.code = code
        self.headers = headers
        self.rows = rows
        self.max_col = max_col
        self.dirty = False

    def row_count(self):
        return len(self.rows)


class ExcelHandler:
    """Excel 文件处理器"""

    def __init__(self, filepath):
        self.filepath = filepath
        self.wb = None
        self.sheet_names = []

    def load(self):
        """加载工作簿，返回 sheet_names 列表"""
        self.wb = openpyxl.load_workbook(self.filepath, data_only=False)
        self.sheet_names = self.wb.sheetnames
        return self.sheet_names

    def _detect_effective_cols(self, ws):
        """检测有效列数：从前几行向右扫描，找到最后一个非空列"""
        max_col = 1
        scan_rows = min(ws.max_row, DATA_START_ROW)
        for row_idx in range(1, scan_rows + 1):
            for col_idx in range(1, ws.max_column + 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None and str(val).strip() != "":
                    if col_idx > max_col:
                        max_col = col_idx
        return max_col

    def read_sheet(self, sheet_name):
        """读取单个工作簿，返回 SheetData"""
        ws = self.wb[sheet_name]
        max_col = self._detect_effective_cols(ws)

        title = ws.cell(row=TITLE_ROW, column=1).value or ""
        code = ws.cell(row=CODE_ROW, column=1).value or ""

        headers = []
        for c in range(1, max_col + 1):
            val = ws.cell(row=HEADER_ROW, column=c).value
            headers.append(str(val).strip() if val is not None else f"列{c}")

        rows = []
        for r in range(DATA_START_ROW, ws.max_row + 1):
            row_data = []
            has_data = False
            for c in range(1, max_col + 1):
                val = ws.cell(row=r, column=c).value
                if val is not None and str(val).strip() != "":
                    has_data = True
                row_data.append(val)
            if has_data:
                rows.append(row_data)

        return SheetData(sheet_name, title, code, headers, rows, max_col)

    def write_sheet(self, sheet_data):
        """将修改后的数据写回工作簿（内存中，未保存到磁盘）"""
        ws = self.wb[sheet_data.name]
        max_col = sheet_data.max_col

        ws.cell(row=TITLE_ROW, column=1, value=sheet_data.title)
        ws.cell(row=CODE_ROW, column=1, value=sheet_data.code)

        for c in range(1, max_col + 1):
            ws.cell(row=HEADER_ROW, column=c, value=sheet_data.headers[c - 1])

        for r in range(DATA_START_ROW, ws.max_row + 1):
            for c in range(1, max_col + 1):
                ws.cell(row=r, column=c, value=None)

        for i, row_data in enumerate(sheet_data.rows):
            r = DATA_START_ROW + i
            for c in range(1, max_col + 1):
                val = row_data[c - 1] if c - 1 < len(row_data) else None
                ws.cell(row=r, column=c, value=val)

        sheet_data.dirty = False

    def save(self, path=None):
        """保存工作簿到磁盘"""
        save_path = path or self.filepath
        self.wb.save(save_path)
        if path:
            self.filepath = path

    def get_sheet_info(self):
        """返回各工作簿的概要信息（名称、行数、列数）"""
        info = []
        for name in self.sheet_names:
            ws = self.wb[name]
            max_col = self._detect_effective_cols(ws)
            data_rows = max(0, ws.max_row - DATA_START_ROW + 1)
            info.append((name, data_rows, max_col))
        return info
